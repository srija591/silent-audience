import re
from pathlib import Path
from .speech import words

STOP = set("a an the and or but of to in on for with at by from is are was were be been this that these those it its as we you they i our your their can will would should could not no yes so if then than into over about more most also just very have has had do does did".split())


def parse_slides(path: Path):
    ext = path.suffix.lower()
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for s in Presentation(str(path)).slides:
            parts = []
            for sh in s.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
            out.append("\n".join(parts).strip())
        return out
    if ext == ".pdf":
        import fitz
        with fitz.open(str(path)) as d:
            return [p.get_text().strip() for p in d]
    return []


def _syll(w):
    w = re.sub(r"e$", "", w.lower())
    return max(1, len(re.findall(r"[aeiouy]+", w)))


def quality(slides):
    per, all_w = [], []
    for i, t in enumerate(slides, 1):
        w = words(t)
        all_w += w
        per.append({"slide": i, "words": len(w), "title": (t.splitlines() or [""])[0][:60]})
    sents = max(1, sum(len(re.findall(r"[.!?\n]", t)) for t in slides))
    syl = sum(_syll(w) for w in all_w) or 1
    nw = max(1, len(all_w))
    flesch = round(206.835 - 1.015 * (nw / sents) - 84.6 * (syl / nw), 1)
    jargon = round(100 * sum(1 for w in all_w if _syll(w) >= 4 or len(w) >= 12) / nw, 1)
    avg = round(nw / max(1, len(slides)))
    return {"slide_count": len(slides), "avg_words_per_slide": avg, "readability": max(0, min(100, flesch)),
            "jargon_density_pct": jargon, "per_slide": per, "dense_slides": [p["slide"] for p in per if p["words"] > 80]}


def _terms(t):
    return {w for w in words(t) if w not in STOP and len(w) > 3}


def align(entries, slides):
    """Assign each speech chunk to a slide (monotone, keyword overlap). Returns per-slide time, speaker, text."""
    terms = [_terms(s) for s in slides]
    cur = 0
    res = [{"slide": i + 1, "seconds": 0.0, "speakers": {}, "text": ""} for i in range(len(slides))]
    for e in entries:
        ew = _terms(e["text"])
        best, score = cur, 0
        for j in range(cur, min(len(slides), cur + 3)):
            sc = len(ew & terms[j])
            if sc > score:
                best, score = j, sc
        cur = best
        r = res[cur]
        d = max(0.0, e["end"] - e["start"])
        r["seconds"] += d
        r["speakers"][e["speaker"]] = r["speakers"].get(e["speaker"], 0) + d
        r["text"] += " " + e["text"]
    return res


def heuristic_coverage(aligned, slides):
    out = []
    for a, s in zip(aligned, slides):
        st, sp = _terms(s), set(words(a["text"]))
        cov = len(st & sp) / len(st) if st else 1
        sw, spw = words(s), words(a["text"])
        grams = {tuple(sw[i:i + 4]) for i in range(len(sw) - 3)}
        sp_grams = {tuple(spw[i:i + 4]) for i in range(len(spw) - 3)}
        verb = len(grams & sp_grams) / len(grams) if grams else 0
        title = (s.splitlines() or [f"Slide {a['slide']}"])[0][:50]
        if a["seconds"] < 3 or cov < 0.15:
            status, note = "skipped", "Barely discussed in the talk."
        elif verb > 0.5:
            status, note = "verbatim", "Slide text was read almost word for word."
        elif cov < 0.4:
            status, note = "insufficient", "Only a small part of this slide's content was explained."
        else:
            status, note = "strong", "Good alignment between what was said and the slide."
        out.append({"slide": a["slide"], "title": title, "status": status, "note": note})
    return out
